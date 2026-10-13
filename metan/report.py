"""Автоматический отчёт в Word (.docx) для НИР."""

import datetime
import io

import matplotlib
from matplotlib.figure import Figure

from . import __version__
from . import plotting
from .analysis import exit_summary
from .kernels import NF, NT
from .params import Params

matplotlib.rcParams.setdefault("font.family", "DejaVu Sans")


def _fig_png(draw, size=(7.0, 3.6)):
    fig = Figure(figsize=size, dpi=200)
    ax = fig.add_subplot(111)
    draw(fig, ax)
    plotting.style(fig, "light")
    fig.tight_layout()
    buf = io.BytesIO()
    fig.savefig(buf, format="png", dpi=200)
    buf.seek(0)
    return buf


class _Doc:
    def __init__(self):
        from docx import Document
        from docx.shared import Pt
        self.doc = Document()
        st = self.doc.styles["Normal"]
        st.font.name = "Times New Roman"
        st.font.size = Pt(12)
        self.nfig = 0
        self.ntab = 0

    def h(self, text, level=1):
        self.doc.add_heading(text, level=level)

    def p(self, text, italic=False):
        par = self.doc.add_paragraph()
        run = par.add_run(text)
        run.italic = italic
        return par

    def figure(self, png, caption, width_cm=16):
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.shared import Cm
        self.doc.add_picture(png, width=Cm(width_cm))
        self.doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
        self.nfig += 1
        par = self.doc.add_paragraph("Рисунок %d – %s" % (self.nfig, caption))
        par.alignment = WD_ALIGN_PARAGRAPH.CENTER

    def table(self, caption, header, rows):
        from docx.shared import Pt
        self.ntab += 1
        self.doc.add_paragraph("Таблица %d – %s" % (self.ntab, caption))
        t = self.doc.add_table(rows=1, cols=len(header))
        t.style = "Table Grid"
        for c, text in enumerate(header):
            cell = t.rows[0].cells[c]
            cell.text = str(text)
            for r in cell.paragraphs[0].runs:
                r.bold = True
                r.font.size = Pt(10)
        for row in rows:
            cells = t.add_row().cells
            for c, text in enumerate(row):
                cells[c].text = str(text)
                for r in cells[c].paragraphs[0].runs:
                    r.font.size = Pt(10)
        self.doc.add_paragraph()


def _fmt(x, nd=4):
    try:
        x = float(x)
    except (TypeError, ValueError):
        return str(x)
    if x != x:
        return "—"
    if abs(x) >= 1e4 or (abs(x) < 1e-3 and x != 0):
        return "%.*e" % (nd - 1, x)
    return "%.*g" % (nd + 1, x)


def build(path, results, series=None, opt=None, gridcheck=None, title=None):
    """Собрать отчёт. results - Results; series/opt/gridcheck - необязательно."""
    d = _Doc()
    p = results.params
    d.doc.add_heading(title or "Расчёт горения метана в камере сгорания", 0)
    d.p("Дата: %s. Программа Metan %s (Python-версия met21)."
        % (datetime.datetime.now().strftime("%d.%m.%Y %H:%M"), __version__),
        italic=True)

    d.h("1. Исходные данные")
    rows = []
    for name, _, meta in Params.field_info():
        if name in ("out10", "out20"):
            continue
        rows.append((meta["label"], name, getattr(p, name), meta["unit"]))
    d.table("Исходные данные расчёта", ["Параметр", "Обозначение",
                                         "Значение", "Ед."], rows)

    v0 = results.variants[0]
    d.figure(_fig_png(lambda f, a: plotting.draw_geometry(a, v0)),
             "Схема камеры и расчётная сетка (размеры в мм)")

    d.h("2. Результаты по вариантам")
    rows = []
    for v in results.variants:
        ex = exit_summary(v)
        val = v.values
        rows.append((v.NVAR, "%.2f" % val["ALFA"], "%.2f" % val["KM"],
                     _fmt(val["Ms"]), "%.1f" % val["Vgor"], "%.1f" % val["Vok"],
                     v.NITER, "%.5f (%s)" % (v.fkomp, v.fkomp_kind.lower()),
                     "%.0f" % ex["T"],
                     "расходимость" if v.diverged else "сошёлся"))
    d.table("Основные результаты",
            ["№", "α", "KM", "Ms, кг/с", "Vгор, м/с", "Vок, м/с", "NITER",
             "Fkomp", "T вых, K", "Статус"], rows)
    if len(results.variants) > 1:
        xs = [v.values["ALFA"] for v in results.variants]
        ys = [v.fkomp for v in results.variants]
        d.figure(_fig_png(lambda f, a: plotting.draw_curves(
            a, [("", xs, ys)], "α", "Fkomp"), size=(6, 3)),
            "Коэффициент полноты сгорания в зависимости от α")
    d.figure(_fig_png(lambda f, a: plotting.draw_convergence(a, results.variants[:3],
                                                             keys=("psi", "T"))),
             "Сходимость итераций (изменение ψ и T за итерацию)")

    d.h("3. Поля течения")
    for v in results.variants:
        a = v.values["ALFA"]
        d.figure(_fig_png(lambda f, ax, v=v: plotting.draw(
            ax, v, NT, colorbar_cb=lambda m, l: f.colorbar(m, ax=ax, label="T, K",
                                                          shrink=0.8))),
            "Поле температуры, α = %.2f" % a)
        d.figure(_fig_png(lambda f, ax, v=v: plotting.draw(
            ax, v, NF, style="lines", mirror=False, labels=True, units="m",
            stream=False)),
            "Распределение функции тока ψ по объёму КС, α = %.2f" % a)

    n = 4
    if series is not None:
        from .params import Params as _P
        meta = {k: m for k, _, m in _P.field_info()}[series.name]
        d.h("%d. Серия расчётов по параметру %s" % (n, series.name))
        n += 1
        alphas = series.alphas()
        curves = []
        for i, a in enumerate(alphas):
            x, y = series.curve("fkomp", i + 1)
            curves.append(("α=%.2f" % a, x, y))
        d.figure(_fig_png(lambda f, ax: plotting.draw_curves(
            ax, curves, "%s, %s" % (meta["label"], meta["unit"]), "Fkomp")),
            "Fkomp в зависимости от параметра %s" % series.name)
        rows = []
        for val, s, err in series.rows:
            rows.append([_fmt(val)] + ([("%.5f" % r["fkomp"]) if not r["diverged"]
                                        else "расх." for r in s] if s else
                                       [err] * max(1, len(alphas))))
        d.table("Fkomp по серии %s" % series.name,
                [series.name] + ["α=%.2f" % a for a in alphas] if alphas
                else [series.name, "ошибка"], rows)
    if opt is not None:
        d.h("%d. Подбор параметров на максимум Fkomp" % n)
        n += 1
        d.p("α = %.2f. Начальное значение Fkomp = %.5f, после подбора %.5f "
            "(%d расчётов)." % (opt.alpha, opt.start_f, opt.best_f, len(opt.evals)))
        d.table("Результат подбора", ["Параметр", "Было", "Стало"],
                [(nm, _fmt(a), _fmt(b)) for nm, a, b in
                 zip(opt.names, opt.start_x, opt.best_x)])
    if gridcheck is not None:
        d.h("%d. Проверка сеточной сходимости" % n)
        gb, gf = gridcheck.base_params, gridcheck.fine_params
        d.p("Базовая сетка %d×%d узлов, мелкая %d×%d (каждый интервал разбит "
            "пополам, итераций %d)." % (gb.IN, gb.JN, gf.IN, gf.JN, gf.NMAX))
        d.table("Сравнение результатов на двух сетках",
                ["α", "Величина", "Базовая", "Мелкая", "Отличие, %"],
                [("%.2f" % a, nm, _fmt(x), _fmt(y),
                  "—" if rel != rel else "%.2f" % rel + (" (расх.)" if bad else ""))
                 for a, nm, x, y, rel, bad in gridcheck.table()])
    d.doc.save(path)
    return path
